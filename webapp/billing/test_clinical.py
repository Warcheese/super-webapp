from io import BytesIO
from pathlib import Path
import tempfile
from unittest.mock import patch
from zipfile import ZipFile
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings, Client
from django.urls import reverse
from docx import Document
from .models import Patient, TreatmentNote, TreatmentAddendum, ClinicalAttachment


class ClinicalTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        config = override_settings(DATA_DIR=Path(self.tmp.name), MEDIA_ROOT=Path(self.tmp.name))
        config.enable()
        self.addCleanup(config.disable)
        self.user = get_user_model().objects.create_user('steffi', password='test-password')
        self.client.force_login(self.user)
        self.patient = Patient.objects.create(first_name='Max', last_name='Müller')
        self.other = Patient.objects.create(first_name='Eva', last_name='Test')
        self.note = TreatmentNote.objects.create(patient=self.patient)
        self.edit = reverse('note_edit', args=[self.patient.pk, self.note.pk])

    def data(self, **changes):
        result = {'date': '2026-10-10', 'body': 'Symptome & Befund: Schulter\nBehandlung dokumentiert.', 'revision': self.note.revision, 'action': 'save'}
        result.update(changes)
        return result

    def finalize(self):
        response = self.client.post(self.edit, self.data(action='finalize', confirm='yes'))
        self.assertEqual(response.status_code, 302)
        self.note.refresh_from_db()

    def docx(self):
        doc = Document()
        doc.add_paragraph('Originaldokument mit ä, ö und ü.')
        stream = BytesIO()
        doc.save(stream)
        return stream.getvalue()

    def upload(self, data=None, name='Altakte.docx'):
        return self.client.post(reverse('attachment_upload', args=[self.patient.pk]), {'file': SimpleUploadedFile(name, self.docx() if data is None else data)})

    def test_pages_render_and_authentication(self):
        urls = [reverse('clinical_record', args=[self.patient.pk]), self.edit, reverse('documentation_export', args=[self.patient.pk])]
        for url in urls:
            self.assertEqual(self.client.get(url).status_code, 200)
        self.client.logout()
        for url in urls:
            self.assertEqual(self.client.get(url).status_code, 302)

    def test_new_note_post_only_and_author(self):
        url = reverse('note_new', args=[self.patient.pk])
        self.assertEqual(self.client.get(url).status_code, 405)
        self.assertEqual(self.client.post(url).status_code, 302)
        self.assertEqual(TreatmentNote.objects.latest('pk').created_by, self.user)

    def test_autosave_and_stale_edit(self):
        response = self.client.post(self.edit, self.data(), HTTP_ACCEPT='application/json')
        self.assertEqual(response.json()['revision'], 1)
        response = self.client.post(self.edit, self.data(body='Überholt'), HTTP_ACCEPT='application/json')
        self.assertEqual(response.status_code, 409)
        self.note.refresh_from_db()
        self.assertIn('Schulter', self.note.body)
        self.assertEqual(self.note.edited_by, self.user)

    def test_invalid_date_not_saved(self):
        response = self.client.post(self.edit, self.data(date='keine Angabe'), HTTP_ACCEPT='application/json')
        self.assertEqual(response.status_code, 400)
        self.note.refresh_from_db()
        self.assertEqual(self.note.revision, 0)

    def test_finalize_requires_text_and_confirmation(self):
        for changes in [dict(action='finalize'), dict(action='finalize', body='', confirm='yes')]:
            self.client.post(self.edit, self.data(**changes))
            self.note.refresh_from_db()
            self.assertIsNone(self.note.finalized_at)

    def test_finalized_note_is_immutable_and_snapshots_identity(self):
        self.finalize()
        self.assertEqual(self.note.author_name, 'steffi')
        self.assertEqual(self.note.patient_snapshot['last_name'], 'Müller')
        response = self.client.post(self.edit, self.data(body='Änderung'), HTTP_ACCEPT='application/json')
        self.assertEqual(response.status_code, 409)
        self.note.body = 'Änderung'
        with self.assertRaises(ValidationError):
            self.note.save()
        self.assertContains(self.client.get(reverse('note_detail', args=[self.patient.pk, self.note.pk])), 'Schulter')

    def test_addendum_confirmed_and_original_preserved(self):
        url = reverse('note_addendum', args=[self.patient.pk, self.note.pk])
        self.assertEqual(self.client.post(url, {'body': 'Nachtrag', 'confirm': 'on'}).status_code, 404)
        self.finalize()
        original = self.note.body
        self.assertEqual(self.client.post(url, {'body': 'Nachtrag'}).status_code, 400)
        self.assertEqual(self.client.post(url, {'body': 'Zusatzbefund', 'confirm': 'on'}).status_code, 302)
        addition = self.note.addenda.get()
        self.assertEqual(addition.author_name, 'steffi')
        with self.assertRaises(ValidationError):
            addition.save()
        self.note.refresh_from_db()
        self.assertEqual(self.note.body, original)
        self.assertContains(self.client.get(reverse('clinical_record', args=[self.patient.pk])+'?q=Zusatzbefund'), 'Schulter')

    def test_patient_documentation_protects_deletion(self):
        self.client.post(reverse('patient_delete', args=[self.patient.pk]))
        self.assertTrue(Patient.objects.filter(pk=self.patient.pk).exists())
        self.patient.active = False
        self.patient.save()
        self.assertEqual(self.client.get(reverse('clinical_record', args=[self.patient.pk])).status_code, 200)

    def test_upload_original_download_integrity(self):
        content = self.docx()
        self.assertEqual(self.upload(content).status_code, 302)
        attachment = ClinicalAttachment.objects.get()
        self.assertTrue(attachment.file.name.startswith('documents/'))
        self.assertEqual(Path(attachment.file.path).read_bytes(), content)
        url = reverse('attachment_download', args=[self.patient.pk, attachment.pk])
        response = self.client.get(url)
        self.assertEqual(b''.join(response.streaming_content), content)
        self.assertIn('no-store', response['Cache-Control'])
        Path(attachment.file.path).write_bytes(b'changed')
        self.assertEqual(self.client.get(url).status_code, 409)

    def test_invalid_uploads_rejected(self):
        for name, content in [('test.docx', b'broken'), ('test.pdf', self.docx())]:
            self.assertEqual(self.upload(content, name).status_code, 400)
        stream = BytesIO()
        with ZipFile(stream, 'w') as archive:
            archive.writestr('[Content_Types].xml', '')
            archive.writestr('word/document.xml', '')
            archive.writestr('word/vbaProject.bin', 'macro')
        self.assertEqual(self.upload(stream.getvalue()).status_code, 400)
        self.assertFalse(ClinicalAttachment.objects.exists())

    def test_wrong_patient_cannot_access_notes_or_files(self):
        self.upload()
        attachment = ClinicalAttachment.objects.get()
        for route, pk in [('note_edit', self.note.pk), ('note_detail', self.note.pk), ('note_export', self.note.pk), ('attachment_download', attachment.pk)]:
            self.assertEqual(self.client.get(reverse(route, args=[self.other.pk, pk])).status_code, 404)

    def test_word_export_draft_final_addenda_and_original_listing(self):
        url = reverse('note_export', args=[self.patient.pk, self.note.pk])
        response = self.client.get(url)
        paragraphs = Document(BytesIO(b''.join(response.streaming_content))).paragraphs
        self.assertIn('ENTWURF', '\n'.join(p.text for p in paragraphs))
        self.finalize()
        TreatmentAddendum.objects.create(note=self.note, body='Nachtrag: Wärme', author_name='steffi')
        self.patient.last_name = 'Geändert'
        self.patient.save()
        response = self.client.get(url)
        text = '\n'.join(p.text for p in Document(BytesIO(b''.join(response.streaming_content))).paragraphs)
        self.assertIn('Müller', text)
        self.assertIn('Nachtrag: Wärme', text)
        self.assertNotIn('ENTWURF', text)
        self.upload()
        response = self.client.get(reverse('documentation_export', args=[self.patient.pk]))
        text = '\n'.join(p.text for p in Document(BytesIO(b''.join(response.streaming_content))).paragraphs)
        self.assertIn('Altakte.docx', text)
        self.assertNotIn('Originaldokument mit', text)

    def test_csrf_required_for_mutations(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.user)
        self.assertEqual(client.post(self.edit, self.data()).status_code, 403)

    def test_invalid_control_characters_rejected(self):
        response = self.client.post(self.edit, self.data(body='Befund\x01'), HTTP_ACCEPT='application/json')
        self.assertEqual(response.status_code, 400)
        self.note.refresh_from_db()
        self.assertEqual(self.note.body, '')

    def test_failed_database_upload_removes_file(self):
        with patch('billing.clinical_views.ClinicalAttachment.save', side_effect=RuntimeError('synthetic failure')):
            with self.assertRaises(RuntimeError):
                self.upload()
        self.assertEqual(list(Path(self.tmp.name).rglob('*.docx')), [])
        self.assertFalse(ClinicalAttachment.objects.exists())

    def test_attachment_path_escape_and_authentication(self):
        self.upload()
        attachment = ClinicalAttachment.objects.get()
        url = reverse('attachment_download', args=[self.patient.pk, attachment.pk])
        self.client.logout()
        self.assertEqual(self.client.get(url).status_code, 302)
        self.client.force_login(self.user)
        attachment.file.name = '../outside.docx'
        attachment.save()
        self.assertEqual(self.client.get(url).status_code, 404)

    def test_search_attachment_filename_and_note_text(self):
        self.upload()
        self.note.body = 'Kopfschmerz'
        self.note.save()
        url = reverse('clinical_record', args=[self.patient.pk])
        response = self.client.get(url, {'q': 'Altakte'})
        self.assertContains(response, 'Altakte.docx')
        self.assertNotContains(response, 'Kopfschmerz')
        response = self.client.get(url, {'q': 'Kopf'})
        self.assertContains(response, 'Kopfschmerz')
        self.assertNotContains(response, 'Altakte.docx')
