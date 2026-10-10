from io import BytesIO
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from django.utils import timezone


def export_documentation(patient, notes, attachments=()):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    normal = document.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    for name in ['Title', 'Heading 1', 'Heading 2']:
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    document.add_heading('Behandlungsdokumentation', 0)
    notes = list(notes)
    identity = notes[0].patient_snapshot if len(notes) == 1 and notes[0].finalized_at else patient.snapshot()
    document.add_paragraph(f"Patient: {identity['last_name']}, {identity['first_name']}")
    if identity.get('birth_date'):
        from datetime import date
        document.add_paragraph('Geburtsdatum: '+date.fromisoformat(identity['birth_date']).strftime('%d.%m.%Y'))
    document.add_paragraph('Dokumentation der Behandlungen und zugehöriger Nachträge. Entwürfe sind ausdrücklich gekennzeichnet.')
    for note in notes:
        document.add_heading(note.date.strftime('%d.%m.%Y')+(' ENTWURF' if not note.finalized_at else ''), 1)
        if note.finalized_at:
            document.add_paragraph(f"Abgeschlossen am {timezone.localtime(note.finalized_at):%d.%m.%Y %H:%M} von {note.author_name}")
        else:
            document.add_paragraph(f"Entwurf zuletzt gespeichert am {timezone.localtime(note.updated_at):%d.%m.%Y %H:%M}")
        for paragraph in (note.body or 'Noch kein Text vorhanden.').splitlines():
            document.add_paragraph(paragraph)
        for addendum in note.addenda.all():
            document.add_heading('Nachtrag', 2)
            document.add_paragraph(f"{timezone.localtime(addendum.created_at):%d.%m.%Y %H:%M} · {addendum.author_name}")
            for paragraph in addendum.body.splitlines():
                document.add_paragraph(paragraph)
    if not notes:
        document.add_paragraph('Noch keine Behandlungen dokumentiert.')
    attachments = list(attachments)
    if attachments:
        document.add_heading('Angehängte Originaldokumente', 1)
        document.add_paragraph('Die Dateien sind separat in der Patientenakte verfügbar. Ihr Inhalt wird nicht in diesen Auszug übernommen.')
        for attachment in attachments:
            document.add_paragraph(f"{attachment.original_name} · hochgeladen am {timezone.localtime(attachment.uploaded_at):%d.%m.%Y %H:%M}")
    document.core_properties.title = 'Behandlungsdokumentation'
    document.core_properties.author = ''
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
